from __future__ import annotations

import base64
import csv
import io
import re
import zipfile
from dataclasses import dataclass
from xml.parsers import expat

from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string, range_boundaries

MAX_UPLOAD_BYTES = 8 * 1024 * 1024
MAX_QUESTIONS = 500
MAX_QUESTION_CHARS = 4000
MAX_SCAN_ROWS = 5000
MAX_COLUMNS = 64
MAX_CELLS = 100_000
MAX_CELL_CHARS = 32_768
MAX_WORKSHEETS = 20
MAX_ZIP_MEMBERS = 256
MAX_ZIP_ENTRY_BYTES = 16 * 1024 * 1024
MAX_XLSX_EXPANDED_BYTES = 32 * 1024 * 1024
MAX_XLSX_XML_ELEMENTS = 500_000
QUESTION_HEADERS = {"question", "questions", "query", "prompt", "คำถาม", "โจทย์"}


class QuestionImportLimitError(ValueError):
    """The input exceeded a limit; never silently try another worksheet."""


@dataclass(frozen=True)
class ImportedQuestions:
    filename: str
    sheet: str | None
    header: str
    questions: list[str]
    row_numbers: list[int]


def _clean(value) -> str:
    if value is None:
        return ""
    return str(value).replace("\ufeff", "").strip()


def _normalize_header(value) -> str:
    return " ".join(_clean(value).casefold().split())


def _find_question_column(rows: list[list[object]]) -> tuple[int, int, str, bool]:
    """Return (header_row_index, column_index, display_header, has_real_header)."""
    for ridx, row in enumerate(rows[:20]):
        for cidx, value in enumerate(row):
            normalized = _normalize_header(value)
            if normalized in QUESTION_HEADERS:
                return ridx, cidx, _clean(value) or "question", True

    # Organizer files sometimes contain a blank/odd header. Pick the column
    # containing the most substantial text and treat row 1 as data.
    if rows:
        width = max((len(r) for r in rows[:50]), default=0)
        best_col, best_score = -1, -1
        for cidx in range(width):
            score = 0
            for row in rows[:50]:
                value = _clean(row[cidx] if cidx < len(row) else "")
                if len(value) >= 8:
                    score += 1
            if score > best_score:
                best_col, best_score = cidx, score
        if best_col >= 0 and best_score > 0:
            return -1, best_col, "question", False

    raise ValueError("ไม่พบคอลัมน์คำถาม กรุณาใช้หัวคอลัมน์ question, query, prompt, คำถาม หรือ โจทย์")


def _extract(rows: list[list[object]], filename: str, sheet: str | None) -> ImportedQuestions:
    if not rows:
        raise ValueError("ไฟล์ไม่มีข้อมูล")

    header_row, question_col, header, has_real_header = _find_question_column(rows)
    questions: list[str] = []
    row_numbers: list[int] = []
    start = header_row + 1 if has_real_header else 0

    for ridx in range(start, len(rows)):
        row = rows[ridx]
        value = _clean(row[question_col] if question_col < len(row) else "")
        if not value:
            continue
        # Do not accidentally import a header repeated later in the file.
        if _normalize_header(value) in QUESTION_HEADERS:
            continue
        if len(value) > MAX_QUESTION_CHARS:
            raise QuestionImportLimitError(
                f"Question at row {ridx + 1} exceeds {MAX_QUESTION_CHARS} characters"
            )
        if len(questions) >= MAX_QUESTIONS:
            raise QuestionImportLimitError(f"A file may contain at most {MAX_QUESTIONS} questions")
        questions.append(value)
        row_numbers.append(ridx + 1)

    if not questions:
        raise ValueError("ไม่พบคำถามที่ไม่ว่างในไฟล์")

    return ImportedQuestions(filename=filename, sheet=sheet, header=header, questions=questions, row_numbers=row_numbers)


def _bounded_rows(rows):
    """Bound materialized values even when CSV or worksheet metadata is false."""
    cells = 0
    for index, row in enumerate(rows, 1):
        if index > MAX_SCAN_ROWS:
            raise QuestionImportLimitError(f"A sheet may contain at most {MAX_SCAN_ROWS} rows")
        if len(row) > MAX_COLUMNS:
            raise QuestionImportLimitError(f"A sheet may contain at most {MAX_COLUMNS} columns")
        cells += len(row)
        if cells > MAX_CELLS:
            raise QuestionImportLimitError(f"A sheet may contain at most {MAX_CELLS} cells")
        if any(isinstance(value, str) and len(value) > MAX_CELL_CHARS for value in row):
            raise QuestionImportLimitError(f"A cell may contain at most {MAX_CELL_CHARS} characters")
        yield list(row)


def _check_upload_bytes(data: bytes) -> None:
    if not data:
        raise ValueError("ไฟล์ว่าง")
    if len(data) > MAX_UPLOAD_BYTES:
        raise QuestionImportLimitError("ไฟล์ใหญ่เกิน 8 MB")


def parse_csv_bytes(data: bytes, filename: str) -> ImportedQuestions:
    _check_upload_bytes(data)
    text = None
    last_error = None
    for encoding in ("utf-8-sig", "utf-8", "cp874", "tis-620"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError as exc:
            last_error = exc
    if text is None:
        raise ValueError(f"อ่าน encoding ของ CSV ไม่ได้: {last_error}")

    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    try:
        rows = list(_bounded_rows(csv.reader(io.StringIO(text), dialect, strict=True)))
    except csv.Error as exc:
        raise ValueError("ไฟล์ CSV ไม่ถูกต้องหรือมีช่องข้อมูลใหญ่เกินกำหนด") from exc
    return _extract(rows, filename, None)


def _validate_xlsx_archive(data: bytes) -> None:
    """Check expansion and real XML coordinates before openpyxl allocates cells.

    Worksheet dimensions are user-controlled. A tiny ZIP can otherwise ask
    openpyxl to synthesize billions of blank cells. XML is scanned without
    building a tree, and DTD/entity declarations are never accepted.
    """
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        members = archive.infolist()
        if len(members) > MAX_ZIP_MEMBERS:
            raise QuestionImportLimitError(f"Workbook exceeds {MAX_ZIP_MEMBERS} ZIP members")
        if len({member.filename for member in members}) != len(members):
            raise ValueError("Workbook contains duplicate ZIP members")
        if sum(member.file_size for member in members) > MAX_XLSX_EXPANDED_BYTES:
            raise QuestionImportLimitError("Workbook expanded size exceeds 32 MB")
        sheets = xml_elements = 0
        for member in members:
            if member.file_size > MAX_ZIP_ENTRY_BYTES:
                raise QuestionImportLimitError("Workbook ZIP entry expanded size exceeds 16 MB")
            if member.flag_bits & 1 or member.compress_type not in {zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED}:
                raise ValueError("Workbook contains unsupported ZIP entries")
            if not member.filename.casefold().endswith((".xml", ".rels")):
                continue
            parser = expat.ParserCreate(namespace_separator="}")
            root_seen = False
            is_sheet = False
            row_index = row_count = column_index = cell_count = 0
            max_row = max_column = 0

            def reject_dtd(*_):
                raise ValueError("Workbook XML must not contain DTD or entity declarations")

            def start_element(tag, attrs):
                nonlocal root_seen, is_sheet, sheets, xml_elements
                nonlocal row_index, row_count, column_index, cell_count, max_row, max_column
                xml_elements += 1
                if xml_elements > MAX_XLSX_XML_ELEMENTS:
                    raise QuestionImportLimitError("Workbook XML element count exceeds its limit")
                name = tag.rsplit("}", 1)[-1]
                if not root_seen:
                    root_seen = True
                    is_sheet = name == "worksheet"
                    if is_sheet:
                        sheets += 1
                        if sheets > MAX_WORKSHEETS:
                            raise QuestionImportLimitError(f"Workbook exceeds {MAX_WORKSHEETS} worksheets")
                if not is_sheet:
                    return
                if name == "dimension":
                    bounds = range_boundaries(attrs.get("ref", ""))
                    if not all(isinstance(value, int) and value > 0 for value in bounds):
                        raise ValueError("Workbook has an invalid worksheet dimension")
                    _, _, columns, rows = bounds
                    if rows > MAX_SCAN_ROWS or columns > MAX_COLUMNS or rows * columns > MAX_CELLS:
                        raise QuestionImportLimitError("Worksheet dimensions exceed row, column or cell limits")
                elif name == "row":
                    row_count += 1
                    row_index = int(attrs.get("r", row_index + 1))
                    column_index = 0
                    if row_count > MAX_SCAN_ROWS or not 1 <= row_index <= MAX_SCAN_ROWS:
                        raise QuestionImportLimitError(f"A sheet may contain at most {MAX_SCAN_ROWS} rows")
                    max_row = max(max_row, row_index)
                elif name == "c":
                    cell_count += 1
                    coordinate = attrs.get("r")
                    cell_row = row_index
                    if coordinate:
                        match = re.fullmatch(r"([A-Za-z]{1,3})([1-9][0-9]{0,6})", coordinate)
                        if not match:
                            raise ValueError("Workbook contains an invalid cell coordinate")
                        column_index = column_index_from_string(match[1])
                        cell_row = int(match[2])
                    else:
                        column_index += 1
                    if cell_count > MAX_CELLS or not 1 <= cell_row <= MAX_SCAN_ROWS or column_index > MAX_COLUMNS:
                        raise QuestionImportLimitError("Worksheet cells exceed row, column or cell limits")
                    max_row = max(max_row, cell_row)
                    max_column = max(max_column, column_index)
                    if max_row * max_column > MAX_CELLS:
                        raise QuestionImportLimitError(f"Worksheet scan exceeds {MAX_CELLS} cells")

            parser.StartElementHandler = start_element
            parser.StartDoctypeDeclHandler = reject_dtd
            parser.EntityDeclHandler = reject_dtd
            parser.ExternalEntityRefHandler = reject_dtd
            total_read = 0
            with archive.open(member) as stream:
                while chunk := stream.read(64 * 1024):
                    total_read += len(chunk)
                    if total_read > MAX_ZIP_ENTRY_BYTES:
                        raise QuestionImportLimitError("Workbook ZIP entry expanded size exceeds 16 MB")
                    parser.Parse(chunk, False)
                parser.Parse(b"", True)


def parse_xlsx_bytes(data: bytes, filename: str) -> ImportedQuestions:
    _check_upload_bytes(data)
    try:
        _validate_xlsx_archive(data)
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
    except QuestionImportLimitError:
        raise
    except Exception as exc:
        raise ValueError("ไฟล์ .xlsx ไม่ถูกต้องหรืออ่านไม่ได้") from exc

    try:
        if len(wb.worksheets) > MAX_WORKSHEETS:
            raise QuestionImportLimitError(f"Workbook exceeds {MAX_WORKSHEETS} worksheets")
        errors: list[str] = []
        for ws in wb.worksheets:
            # Real XML coordinates were bounded above. Ignore small, inaccurate
            # dimensions too, so a valid question is never silently skipped.
            ws.reset_dimensions()
            rows = list(_bounded_rows(ws.iter_rows(values_only=True)))
            try:
                return _extract(rows, filename, ws.title)
            except QuestionImportLimitError:
                raise
            except ValueError as exc:
                errors.append(f"{ws.title}: {exc}")
        raise ValueError("ไม่พบตารางคำถามใน workbook: " + " | ".join(errors[:4]))
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("ไฟล์ .xlsx ไม่ถูกต้องหรืออ่านไม่ได้") from exc
    finally:
        wb.close()


def decode_upload(filename: str, content_base64: str) -> ImportedQuestions:
    if not isinstance(filename, str) or not isinstance(content_base64, str):
        raise ValueError("filename and contentBase64 must be strings")
    name = _clean(filename)
    if not name:
        raise ValueError("filename is required")
    if len(name) > 255:
        raise ValueError("filename is limited to 255 characters")
    if len(content_base64) > ((MAX_UPLOAD_BYTES + 2) // 3) * 4:
        raise QuestionImportLimitError("ไฟล์ใหญ่เกิน 8 MB")
    try:
        data = base64.b64decode(content_base64, validate=True)
    except Exception as exc:
        raise ValueError("ไฟล์อัปโหลดไม่ใช่ Base64 ที่ถูกต้อง") from exc
    _check_upload_bytes(data)

    lower = name.casefold()
    if lower.endswith(".csv"):
        return parse_csv_bytes(data, name)
    if lower.endswith(".xlsx"):
        return parse_xlsx_bytes(data, name)
    raise ValueError("รองรับเฉพาะไฟล์ .xlsx และ .csv")
