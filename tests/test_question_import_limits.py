import base64
import io
import re
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook

from competition_ai import question_import as importer


ROOT = Path(__file__).resolve().parents[1]


def workbook_bytes(questions=("AIT เรียนกี่หน่วยกิต?",), *, extra_sheet=False):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["question", "answer"])
    for question in questions:
        sheet.append([question, ""])
    if extra_sheet:
        other = workbook.create_sheet("Other")
        other.append(["question"])
        other.append(["DSBA เรียนกี่ปี?"])
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def rewrite_zip(data, transform):
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as original, zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as rebuilt:
        for member in original.infolist():
            rebuilt.writestr(member, transform(member.filename, original.read(member.filename)))
    return output.getvalue()


def rewrite_sheet(data, transform):
    return rewrite_zip(data, lambda name, content: transform(content) if name == "xl/worksheets/sheet1.xml" else content)


@pytest.mark.parametrize("filename", ["data/evaluation/Easy.xlsx", "knowledge/question_sets/hard_hard10_blank.csv"])
def test_real_repository_question_files_still_import_all_ten(filename):
    path = ROOT / filename
    result = importer.decode_upload(path.name, base64.b64encode(path.read_bytes()).decode("ascii"))
    assert len(result.questions) == len(result.row_numbers) == 10
    assert all(question.strip() for question in result.questions)


def test_tiny_xlsx_with_billion_cell_dimension_is_rejected_before_openpyxl(monkeypatch):
    data = rewrite_sheet(workbook_bytes(), lambda xml: re.sub(br'<dimension ref="[^"]+"', b'<dimension ref="A1:XFD1048576"', xml))
    assert len(data) < 10_000
    monkeypatch.setattr(importer, "load_workbook", lambda *args, **kwargs: pytest.fail("Unsafe worksheet reached openpyxl"))
    with pytest.raises(importer.QuestionImportLimitError, match="dimensions"):
        importer.parse_xlsx_bytes(data, "dimension.xlsx")


def test_false_small_dimension_does_not_silently_drop_questions():
    data = rewrite_sheet(workbook_bytes(), lambda xml: re.sub(br'<dimension ref="[^"]+"', b'<dimension ref="A1:A1"', xml))
    assert importer.parse_xlsx_bytes(data, "small-dimension.xlsx").questions == ["AIT เรียนกี่หน่วยกิต?"]


@pytest.mark.parametrize("coordinate", [b"A5001", b"BM2"])
def test_real_cell_coordinates_are_checked_even_without_dimension(coordinate):
    def edit(xml):
        xml = re.sub(br'<dimension[^>]*/>', b"", xml)
        return xml.replace(b'r="A2"', b'r="' + coordinate + b'"')
    data = rewrite_sheet(workbook_bytes(), edit)
    with pytest.raises(importer.QuestionImportLimitError, match="cells"):
        importer.parse_xlsx_bytes(data, "actual-cell.xlsx")


def test_expanded_zip_size_is_checked_before_openpyxl(monkeypatch):
    data = workbook_bytes()
    monkeypatch.setattr(importer, "MAX_XLSX_EXPANDED_BYTES", 1024)
    monkeypatch.setattr(importer, "load_workbook", lambda *args, **kwargs: pytest.fail("Expanded ZIP reached openpyxl"))
    with pytest.raises(importer.QuestionImportLimitError, match="expanded size"):
        importer.parse_xlsx_bytes(data, "expanded.xlsx")


def test_each_zip_entry_has_an_expansion_bound(monkeypatch):
    monkeypatch.setattr(importer, "MAX_ZIP_ENTRY_BYTES", 1024)
    with pytest.raises(importer.QuestionImportLimitError, match="entry expanded size"):
        importer.parse_xlsx_bytes(workbook_bytes(), "entry.xlsx")


def test_zip_member_count_is_bounded(monkeypatch):
    monkeypatch.setattr(importer, "MAX_ZIP_MEMBERS", 2)
    with pytest.raises(importer.QuestionImportLimitError, match="ZIP members"):
        importer.parse_xlsx_bytes(workbook_bytes(), "members.xlsx")


def test_xml_metadata_element_count_is_bounded(monkeypatch):
    monkeypatch.setattr(importer, "MAX_XLSX_XML_ELEMENTS", 10)
    with pytest.raises(importer.QuestionImportLimitError, match="XML element count"):
        importer.parse_xlsx_bytes(workbook_bytes(), "metadata.xlsx")


def test_worksheet_count_is_bounded(monkeypatch):
    monkeypatch.setattr(importer, "MAX_WORKSHEETS", 1)
    with pytest.raises(importer.QuestionImportLimitError, match="worksheets"):
        importer.parse_xlsx_bytes(workbook_bytes(extra_sheet=True), "sheets.xlsx")


@pytest.mark.parametrize("entry", ["xl/worksheets/sheet1.xml", "xl/_rels/workbook.xml.rels"])
def test_dtd_entity_declarations_are_rejected_before_openpyxl(monkeypatch, entry):
    def edit(name, content):
        if name == entry:
            content = b'<!DOCTYPE worksheet [<!ENTITY unsafe "payload">]>' + content
        return content
    data = rewrite_zip(workbook_bytes(), edit)
    monkeypatch.setattr(importer, "load_workbook", lambda *args, **kwargs: pytest.fail("DTD reached openpyxl"))
    with pytest.raises(ValueError):
        importer.parse_xlsx_bytes(data, "entity.xlsx")


@pytest.mark.parametrize("kind", ["csv", "xlsx"])
def test_question_count_above_500_is_rejected_instead_of_truncated(kind):
    questions = [f"AIT question number {index}" for index in range(importer.MAX_QUESTIONS + 1)]
    data = ("question\n" + "\n".join(questions)).encode() if kind == "csv" else workbook_bytes(questions, extra_sheet=True)
    parser = importer.parse_csv_bytes if kind == "csv" else importer.parse_xlsx_bytes
    with pytest.raises(importer.QuestionImportLimitError, match="500 questions"):
        parser(data, "many." + kind)


@pytest.mark.parametrize("kind", ["csv", "xlsx"])
def test_long_question_is_rejected_with_row_number(kind):
    question = "x" * (importer.MAX_QUESTION_CHARS + 1)
    data = ("question\n" + question).encode() if kind == "csv" else workbook_bytes([question])
    parser = importer.parse_csv_bytes if kind == "csv" else importer.parse_xlsx_bytes
    with pytest.raises(importer.QuestionImportLimitError, match="row 2.*4000"):
        parser(data, "long." + kind)


def test_exact_500_question_boundary_preserves_every_csv_row():
    questions = [f"AIT question number {index}" for index in range(importer.MAX_QUESTIONS)]
    result = importer.parse_csv_bytes(("question\n" + "\n".join(questions)).encode(), "full.csv")
    assert result.questions == questions
    assert result.row_numbers == list(range(2, importer.MAX_QUESTIONS + 2))


def test_csv_column_count_and_empty_row_scanning_are_bounded():
    wide = ("question," + ",".join(["answer"] * importer.MAX_COLUMNS) + "\n").encode()
    with pytest.raises(importer.QuestionImportLimitError, match="columns"):
        importer.parse_csv_bytes(wide, "wide.csv")
    blank_rows = ("question\n" + "\n" * importer.MAX_SCAN_ROWS).encode()
    with pytest.raises(importer.QuestionImportLimitError, match="rows"):
        importer.parse_csv_bytes(blank_rows, "rows.csv")


def test_csv_cells_and_field_length_are_bounded(monkeypatch):
    monkeypatch.setattr(importer, "MAX_CELLS", 3)
    with pytest.raises(importer.QuestionImportLimitError, match="cells"):
        importer.parse_csv_bytes(b"question,answer\nAIT question,unused\n", "cells.csv")
    monkeypatch.setattr(importer, "MAX_CELLS", 100_000)
    large_field = ("question,answer\nAIT question," + "x" * (importer.MAX_CELL_CHARS + 1)).encode()
    with pytest.raises(importer.QuestionImportLimitError, match="cell may contain"):
        importer.parse_csv_bytes(large_field, "field.csv")


def test_malformed_csv_is_a_safe_validation_error():
    with pytest.raises(ValueError, match="CSV"):
        importer.parse_csv_bytes(b'question\n"unterminated', "invalid.csv")


def test_invalid_xlsx_shared_string_references_are_validation_errors():
    data = rewrite_sheet(workbook_bytes(), lambda xml: re.sub(
        br'<c r="A2"[^>]*>.*?</c>', b'<c r="A2" t="s"><v>99999999</v></c>', xml
    ))
    with pytest.raises(ValueError):
        importer.parse_xlsx_bytes(data, "invalid-string.xlsx")


def test_base64_length_is_rejected_before_decoding(monkeypatch):
    monkeypatch.setattr(importer, "MAX_UPLOAD_BYTES", 3)
    monkeypatch.setattr(importer.base64, "b64decode", lambda *args, **kwargs: pytest.fail("Oversized Base64 was decoded"))
    with pytest.raises(importer.QuestionImportLimitError):
        importer.decode_upload("large.csv", "A" * 8)


def test_upload_metadata_is_not_coerced_to_strings():
    with pytest.raises(ValueError, match="must be strings"):
        importer.decode_upload([], "AAAA")
    with pytest.raises(ValueError, match="255"):
        importer.decode_upload("x" * 256, "AAAA")
